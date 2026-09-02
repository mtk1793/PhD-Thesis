"""
Generate Chapter 1 Word document for CAPSM thesis.
Substantial rewrite with HIL validation as first-class methodology.
"""
import docx
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os

def create_chapter1():
    doc = Document()

    # Set default font
    style = doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(12)

    # Title
    title = doc.add_heading('Chapter 1: Introduction', level=1)

    # 1.1
    doc.add_heading('1.1 Background on Power System Challenges in the Renewable Energy Era', level=2)
    doc.add_paragraph(
        'The global energy landscape is undergoing a profound transformation driven by the '
        'imperative to address climate change, enhance energy security, and meet growing electricity '
        'demand. Traditional power systems, characterized by centralized generation, unidirectional '
        'power flow, and relatively predictable operating conditions, are rapidly evolving into complex '
        'networks with diverse energy resources, bidirectional power flows, and high levels of uncertainty. '
        'This transformation is primarily fueled by the integration of renewable energy sources (RES) and the '
        'proliferation of distributed energy resources (DERs).'
    )
    doc.add_paragraph(
        'The International Energy Agency (IEA) reports that renewable energy capacity has more than '
        'doubled in the past decade, with global renewable electricity capacity expected to increase by over '
        '60% between 2020 and 2026, reaching more than 4,800 GW. This unprecedented growth '
        'presents both opportunities and challenges for power system operation, control, and stability.'
    )

    doc.add_heading('Key Challenges', level=3)
    for challenge in [
        'Intermitency and Variability: Renewable energy production is inherently variable and dependent on weather conditions. Solar generation follows diurnal patterns with rapid fluctuations due to cloud cover, while wind generation can vary significantly on multiple timescales.',
        'Reduced System Inertia: Conventional synchronous generators naturally provide rotational inertia that stabilizes grid frequency during disturbances. Inverter-based resources provide little or no inherent inertia.',
        'Bidirectional Power Flows: Traditional distribution networks were designed for unidirectional power flow. DERs transform these into active systems with bidirectional flows.',
        'Increased Complexity: The proliferation of DERs dramatically increases the number of active elements, creating unprecedented complexity in system operation and control.',
        'Cybersecurity Vulnerabilities: The digitalization of power systems creates new cybersecurity vulnerabilities requiring robust security measures integrated with control systems.',
        'Market and Regulatory Challenges: Existing electricity market structures were designed for conventional power systems with centralized generation.'
    ]:
        doc.add_paragraph(challenge, style='List Bullet')

    # 1.2
    doc.add_heading('1.2 Evolution of FACTS Devices and Their Role in Modern Power Systems', level=2)
    doc.add_paragraph(
        'Flexible AC Transmission System (FACTS) devices have emerged as critical technologies '
        'that enhance the controllability and power transfer capability of transmission systems. FACTS devices, '
        'first conceptualized by the Electric Power Research Institute (EPRI) in the late 1980s, represent a class '
        'of power electronic-based controllers that can rapidly and continuously modify the parameters of the '
        'power system.'
    )
    doc.add_heading('FACTS Device Categories', level=3)
    for cat in [
        'Shunt-connected devices: SVC and STATCOM - voltage support through reactive power control',
        'Series-connected devices: TCSC - modifies effective impedance of transmission lines',
        'Combined series-shunt devices: UPFC - comprehensive control over both voltage and power flow'
    ]:
        doc.add_paragraph(cat, style='List Bullet')

    # 1.3
    doc.add_heading('1.3 Integration of EVs into Smart Grids (V2G/G2V)', level=2)
    doc.add_paragraph(
        'The proliferation of electric vehicles (EVs) presents both challenges and opportunities '
        'for power system operation. With appropriate infrastructure and control strategies, EVs can serve as '
        'flexible resources that enhance grid stability and support renewable energy integration through '
        'bidirectional power flow capabilities.'
    )
    doc.add_heading('Operational Modes', level=3)
    for mode in [
        'Grid-to-Vehicle (G2V): The conventional charging mode where power flows from the grid to the EV battery.',
        'Vehicle-to-Grid (V2G): EVs can discharge power back to the grid, functioning as distributed energy storage.',
        'V2H (Vehicle-to-Home): Using the EV battery to power a home during outages or peak price periods.',
        'V2X (Vehicle-to-Everything): A comprehensive framework encompassing all potential power exchange scenarios.'
    ]:
        doc.add_paragraph(mode, style='List Bullet')

    # 1.4
    doc.add_heading('1.4 Research Motivation Highlighting Increasing DER Integration', level=2)
    doc.add_paragraph(
        'The rapid proliferation of Distributed Energy Resources (DERs) represents one of the most '
        'significant transformations in power system architecture. Global DER capacity is projected '
        'to increase from 132.4 GW in 2017 to over 528.4 GW by 2026, representing a compound annual '
        'growth rate of 16.7%.'
    )

    # 1.5
    doc.add_heading('1.5 Problem Statement Identifying Critical Control Gaps', level=2)
    doc.add_paragraph(
        'Despite significant advancements in power system control and the growing application of '
        'artificial intelligence in grid management, several critical gaps remain in controlling modern '
        'power systems with high DER penetration, FACTS devices, and EV fleets. These gaps '
        'represent the core problems that this research aims to address:'
    )

    gaps = [
        ('Gap 1: Inadequate Control Architectures for Complex, Heterogeneous Systems',
         'Traditional hierarchical control architectures face scalability constraints, inability to handle multi-timescale dynamics, and difficulty incorporating partially controllable resources.'),
        ('Gap 2: Limited Real-Time Decision-Making Under Uncertainty',
         'Traditional deterministic approaches are increasingly inadequate. Existing stochastic approaches face computational barriers to real-time implementation.'),
        ('Gap 3: Insufficient Coordination Between FACTS Devices and DERs',
         'FACTS devices and DERs are typically controlled independently, missing opportunities for synergistic operation.'),
        ('Gap 4: Insufficient Fault Detection and Localization Capabilities',
         'Traditional protection schemes were designed for unidirectional power flow. Modern systems with bidirectional flows create protection challenges.'),
        ('Gap 5: Integration Barriers Between Theoretical Advancements and Practical Implementation',
         'A significant gap exists between theoretical AI advancements and their practical implementation. '
         'Controller-Hardware-in-the-loop (HIL) validation on real-time simulators such as OPAL-RT provides the critical bridge '
         'between simulation and real-world deployment. Without HIL validation, AI-based control approaches '
         'remain in academic environments without transitioning to practical implementation.')
    ]
    for title, desc in gaps:
        h = doc.add_heading(title, level=3)
        doc.add_paragraph(desc)

    # 1.6
    doc.add_heading('1.6 Research Objectives with Specific, Measurable Goals', level=2)
    doc.add_paragraph(
        'This research aims to address the identified control gaps through the development of '
        'advanced AI-driven control frameworks validated through hardware-in-the-loop (HIL) testing '
        'on a 4-core OPAL-RT real-time simulator. The following specific, measurable objectives guide this research:'
    )

    objectives = [
        ('Objective 1: Develop a Novel Brain-Inspired AI Control Architecture',
         'Design a multi-layer neural architecture with cognitive and emotional intelligence. '
         'Demonstrate coordination of at least 1,000 heterogeneous resource units. '
         'Achieve control response times under 100 milliseconds for critical events and under5 seconds for non-critical adjustments.'),
        ('Objective 2: Create Advanced Decision-Making Frameworks for Real-Time Operation',
         'Develop reinforcement learning algorithms optimized for power system control. '
         'Achieve at least25% improvement in operational efficiency. '
         'Demonstrate robust performance across at least500 uncertainty scenarios.'),
        ('Objective 3: Formulate Coordinated Control Strategies for FACTS Devices and EV Fleets',
         'Develop mathematical models for coordinated operation of at least three types of FACTS devices with EV fleets. '
         'Design control algorithms that reduce voltage deviations by at least30%.'),
        ('Objective 4: Enhance Fault Detection and Localization Capabilities',
         'Achieve fault detection accuracy exceeding98%. '
         'Reduce fault localization error to less than50 meters. '
         'Decrease fault classification and localization time to under100 milliseconds.'),
        ('Objective 5: Develop a Comprehensive Testing and Validation Framework',
         'Create a three-stage validation pipeline: (1) offline Python/PYPOWER simulation, '
         '(2) Simulink model preparation, (3) Controller-HIL on 4-core OPAL-RT. '
         'Validate control performance across at least200 distinct operational scenarios. '
         'Achieve TRL 6 by demonstrating successful real-time operation on OPAL-RT hardware.')
    ]
    for title, desc in objectives:
        h = doc.add_heading(title, level=3)
        doc.add_paragraph(desc)

    # 1.7
    doc.add_heading('1.7 Original Contributions Clearly Enumerated', level=2)
    doc.add_paragraph(
        'This research makes several original contributions to the field of power system control, '
        'artificial intelligence applications in energy systems, and the integration of distributed energy '
        'resources, FACTS devices, and electric vehicle fleets. The key contributions are enumerated below:'
    )

    contributions = [
        'Contribution 1: Unified Mathematical Framework for Integrated Resource Modeling',
        'Contribution 2: Brain-Inspired Dual-Process AI Architecture for Power System Control',
        'Contribution 3: Quantum-Inspired Reinforcement Learning for Power System Control',
        'Contribution 4: CNN-LSTM Architecture for Real-Time Stability Assessment',
        'Contribution 5: Multi-Modal Fault Detection and Localization Framework',
        'Contribution 6: Coordinated Control Framework for FACTS Devices and EV Fleets',
        'Contribution 7: Digital Twin Implementation Methodology for Power System Control Testing',
        'Contribution 8: Hardware-in-the-Loop Validation Methodology on 4-Core OPAL-RT',
        'Contribution 9: Transferable Knowledge Representation for Cross-Domain Learning',
    ]
    for c in contributions:
        doc.add_paragraph(c, style='List Number')

    # 1.8
    doc.add_heading('1.8 Thesis Organization', level=2)
    doc.add_paragraph(
        'The remainder of this thesis is organized as follows. Each chapter has been '
        'substantially rewritten to incorporate HIL validation as a core methodological component:'
    )

    chapters = [
        'Chapter 2: Literature Review - Added Section 2.8 on real-time simulation and HIL validation in power systems',
        'Chapter 3: Proposed Framework - Added Section 3.10 on HIL-ready architecture design',
        'Chapter4: Advanced AI Control Architecture - Added Section4.7 on computational implementation for real-time deployment',
        'Chapter 5: AI-Based Optimal Placement - Added Section5.8 on HIL-validated placement results',
        'Chapter7: Fault Detection and Localization - Added Section7.6 on HIL validation of fault detection',
        'Chapter8: EV Integration and DER Management - Added Section8.7 on HIL validation of V2G coordination',
        'Chapter10: Testing and Validation Framework - Major expansion as HIL core chapter (3x length)',
        'Chapter11: Results and Performance Analysis - Split into offline Python results + Controller-HIL results',
        'Chapter13: Conclusions and Future Work - Elevated HIL validation to a stated contribution',
    ]
    for ch in chapters:
        doc.add_paragraph(ch, style='List Bullet')

    # References
    doc.add_heading('References', level=2)
    refs = [
        '[1] Hingorani, N.G. (1988). Power electronics in electric utilities. Proceedings of the IEEE, 76(4), 481-482.',
        '[2] Blaabjerg, F., Yang, Y., Yang, D., & Wang, X. (2017). Distributed power-generation systems and protection. Proceedings of the IEEE, 105(7), 1311-1331.',
        '[3] IEA (2021). Renewables 2021 Analysis and forecast. International Energy Agency.',
        '[4] OPAL-RT Technologies (2024). ePHASORSIM: Real-time phasor-based simulation platform. OPAL-RT documentation.',
        '[5] DESL-EPFL (2018). IEEE 39-bus power system Simulink model for OPAL-RT eMegaSim. GitHub repository.',
        '[6] Karamanakos, P., et al. (2020). Model Predictive Control of Power Electronic Systems. IEEE Open J. Ind. Appl., 1, 95-114.',
    ]
    for r in refs:
        doc.add_paragraph(r, style='List Bullet')

    # Save
    out_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'Chapter1_Introduction.docx')
    doc.save(out_path)
    print(f'Saved: {out_path}')
    return out_path

if __name__ == '__main__':
    create_chapter1()
