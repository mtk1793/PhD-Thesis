#!/usr/bin/env python3
"""Build the final CAPSM PhD thesis Word document from markdown chapters.

Pipeline: markdown dialect -> parsed blocks -> python-docx assembly.
Fonts: Times New Roman body 12pt 1.5 spacing; headings per academic spec.
"""
import os, re, glob, sys
from datetime import date

from docx import Document
from docx.shared import Pt, Cm, Inches, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from PIL import Image

BASE = os.path.join(os.path.dirname(__file__), "..")
CH = os.path.join(BASE, "chapters_md")
FIG = os.path.join(BASE)
EQ = os.path.join(BASE, "eq_cache")
FACTS = os.path.join(BASE, "factsheet", "CAPSM_FACTSHEET.md")
OUT = os.path.join(BASE, "CAPSM_Thesis_AcademicPaper_2026-09-09.docx")

TITLE = ("Cognitive Adaptive Power System Management (CAPSM): A Brain-Inspired "
         "Dual-Process AI Framework for Coordinated Control of FACTS Devices, EV Fleets, "
         "and DERs \u2013 Validated Through Controller-Hardware-in-the-Loop Testing on a "
         "4-Core OPAL-RT Real-Time Simulator")
AUTHOR = "Mahmoud Kiasari"
DEPT = "Department of Electrical and Computer Engineering"
UNIV = "Dalhousie University, Halifax, Nova Scotia, Canada"
DEGREE = "A thesis submitted in partial fulfilment of the requirements for the degree of"
DEGNAME = "Doctor of Philosophy"
DATE_STR = "September 2026"

# ----------------------------------------------------------------- helpers
def set_cell_bg(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)

def add_hyperlink(paragraph, url, text, font_size=10):
    part = paragraph.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hl = OxmlElement("w:hyperlink"); hl.set(qn("r:id"), r_id)
    run = OxmlElement("w:r"); rPr = OxmlElement("w:rPr")
    color = OxmlElement("w:color"); color.set(qn("w:val"), "0563C1"); rPr.append(color)
    u = OxmlElement("w:u"); u.set(qn("w:val"), "single"); rPr.append(u)
    sz = OxmlElement("w:sz"); sz.set(qn("w:val"), str(int(font_size * 2))); rPr.append(sz)
    rf = OxmlElement("w:rFonts"); rf.set(qn("w:ascii"), "Times New Roman"); rf.set(qn("w:hAnsi"), "Times New Roman"); rPr.append(rf)
    run.append(rPr)
    t = OxmlElement("w:t"); t.text = text; t.set(qn("xml:space"), "preserve")
    run.append(t); hl.append(run); paragraph._p.append(hl)

def add_field(paragraph, instr, placeholder=None):
    r1 = OxmlElement("w:r"); fld = OxmlElement("w:fldChar"); fld.set(qn("w:fldCharType"), "begin")
    fld.set(qn("w:dirty"), "true"); r1.append(fld)
    r2 = OxmlElement("w:r"); it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve")
    it.text = f" {instr} "; r2.append(it)
    r3 = OxmlElement("w:r"); sep = OxmlElement("w:fldChar"); sep.set(qn("w:fldCharType"), "separate"); r3.append(sep)
    paragraph._p.append(r1); paragraph._p.append(r2); paragraph._p.append(r3)
    if placeholder:
        r4 = OxmlElement("w:r"); t = OxmlElement("w:t"); t.text = placeholder; r4.append(t)
        paragraph._p.append(r4)
    r5 = OxmlElement("w:r"); end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    r5.append(end); paragraph._p.append(r5)

def style_setup(doc):
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"; st.font.size = Pt(12)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    pf = st.paragraph_format
    pf.line_spacing = 1.5; pf.space_after = Pt(6); pf.space_before = Pt(0)
    for name, size in [("Heading 1", 18), ("Heading 2", 14), ("Heading 3", 12)]:
        h = doc.styles[name]
        h.font.name = "Times New Roman"; h.font.size = Pt(size); h.font.bold = True
        h.font.color.rgb = RGBColor(0, 0, 0)
        h.element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        h.element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        h.paragraph_format.space_before = Pt(12); h.paragraph_format.space_after = Pt(6)
        h.paragraph_format.keep_with_next = True

def para(doc, text, align="justify", size=12, bold=False, italic=False,
         space_after=6, style=None, first_indent=None, hanging=None, line_spacing=1.5):
    p = doc.add_paragraph(style=style)
    if align == "justify": p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    elif align == "center": p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "right": p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    else: p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf = p.paragraph_format
    pf.space_after = Pt(space_after); pf.line_spacing = line_spacing
    if first_indent is not None: pf.first_line_indent = Cm(first_indent)
    if hanging is not None:
        pf.first_line_indent = Cm(-hanging); pf.left_indent = Cm(hanging)
    add_rich_runs(p, text, size=size, bold=bold, italic=italic)
    return p

BOLD_ITAL = re.compile(r"(\*\*\*.+?\*\*\*|\*\*.+?\*\*|\*.+?\*)", re.S)

def add_rich_runs(p, text, size=12, bold=False, italic=False):
    for tok in BOLD_ITAL.split(text):
        if not tok: continue
        b, i, txt = bold, italic, tok
        if tok.startswith("***") and tok.endswith("***") and len(tok) > 6:
            b = i = True; txt = tok[3:-3]
        elif tok.startswith("**") and tok.endswith("**") and len(tok) > 4:
            b = True; txt = tok[2:-2]
        elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
            i = True; txt = tok[1:-1]
        r = p.add_run(txt); r.font.size = Pt(size); r.bold = b; r.italic = i
        r.font.name = "Times New Roman"
    return p

# ----------------------------------------------------------------- markdown parsing
def parse_blocks(md_text):
    """Parse the thesis markdown dialect into structured blocks."""
    lines = md_text.split("\n")
    blocks, i = [], 0
    para_buf = []
    def flush():
        if para_buf:
            text = " ".join(para_buf).strip()
            if text: blocks.append(("para", text))
            para_buf.clear()
    while i < len(lines):
        raw = lines[i]; s = raw.strip()
        if s.startswith("```eq"):
            flush()
            m = re.match(r"```eq\n%%([\d\.]+)\n(.*?)\n```", "\n".join(lines[i:]), re.S)
            if m:
                blocks.append(("eq", m.group(1), m.group(2).strip()))
                consumed = m.group(0).count("\n") + 1
                i += consumed; continue
            i += 1; continue
        if s.startswith("```"):
            flush()
            j = i + 1; code = []
            while j < len(lines) and not lines[j].strip().startswith("```"):
                code.append(lines[j]); j += 1
            blocks.append(("code", "\n".join(code)))
            i = j + 1; continue
        if s.startswith("#"):
            flush(); level = len(s) - len(s.lstrip("#")); text = s.lstrip("#").strip()
            blocks.append(("h%d" % min(level, 4), text))
            i += 1; continue
        if s.startswith("!["):
            flush()
            m = re.match(r"!\[(.*?)\]\((.*?)\)", s)
            if m: blocks.append(("fig", m.group(1), m.group(2)))
            i += 1; continue
        if s.startswith("**Table ") or s.startswith("**Algorithm "):
            flush(); blocks.append(("caption", s)); i += 1; continue
        if s.startswith("|"):
            flush(); rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(lines[i].strip()); i += 1
            blocks.append(("table", rows)); continue
        if s.startswith("- "):
            flush(); blocks.append(("li", s[2:].strip())); i += 1; continue
        if not s:
            flush(); i += 1; continue
        para_buf.append(s); i += 1
    flush()
    return blocks

def make_data_table(doc, rows):
    grid = []
    for r in rows:
        cells = [c.strip() for c in r.strip("|").split("|")]
        if all(re.fullmatch(r":?-{2,}:?", c) for c in cells):  # separator
            continue
        grid.append(cells)
    if not grid: return
    ncols = max(len(r) for r in grid)
    t = doc.add_table(rows=len(grid), cols=ncols)
    t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = True
    for ri, row in enumerate(grid):
        for ci in range(ncols):
            cell = t.cell(ri, ci)
            txt = row[ci] if ci < len(row) else ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2); p.paragraph_format.line_spacing = 1.0
            if ri == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_rich_runs(p, txt.replace("**", "").replace("*", ""), size=9.5)
            for r_ in p.runs: r_.font.name = "Times New Roman"
            if ri == 0:
                set_cell_bg(cell, "F2F2F2")
                for r_ in p.runs: r_.bold = True
    for row in t.rows:
        for cell in row.cells:
            cell.paragraphs[0].paragraph_format.space_before = Pt(2)
    para(doc, "", space_after=4, align="left")

# ----------------------------------------------------------------- doc build
def build():
    doc = Document()
    style_setup(doc)

    # ----- section 1: title page (no header/footer)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(2.5)
    sec.left_margin = sec.right_margin = Cm(2.5)
    sec.header.is_linked_to_previous = False
    sec.footer.is_linked_to_previous = False

    para(doc, "", space_after=30, align="center")
    p = para(doc, TITLE, align="center", size=18, bold=True, space_after=24, line_spacing=1.3)
    para(doc, AUTHOR, align="center", size=14, bold=True, space_after=6)
    para(doc, DEPT, align="center", size=12, space_after=2)
    para(doc, UNIV, align="center", size=12, space_after=24)
    para(doc, DEGREE, align="center", size=12, italic=True, space_after=2)
    para(doc, DEGNAME, align="center", size=13, bold=True, space_after=24)
    para(doc, DATE_STR, align="center", size=12, space_after=18)
    para(doc, "Keywords: cognitive control; dual-process AI; quantum-inspired reinforcement "
              "learning; FACTS devices; vehicle-to-grid; hardware-in-the-loop validation",
         align="center", size=10.5, italic=True, space_after=6)

    # ----- section 2: front matter + TOC (with header/footer)
    sec2 = doc.add_section(WD_SECTION.NEW_PAGE)
    sec2.page_width, sec2.page_height = Cm(21.0), Cm(29.7)
    sec2.top_margin = sec2.bottom_margin = Cm(2.5)
    sec2.left_margin = sec2.right_margin = Cm(2.5)
    sec2.header.is_linked_to_previous = False
    sec2.footer.is_linked_to_previous = False
    hp = sec2.header.paragraphs[0]; hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hr = hp.add_run("CAPSM PhD Thesis"); hr.font.size = Pt(10); hr.font.name = "Times New Roman"
    fp = sec2.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = fp.add_run("Page "); fr.font.size = Pt(10); fr.font.name = "Times New Roman"
    add_field(fp, "PAGE")
    for r in fp.runs: r.font.size = Pt(10)

    fm = open(os.path.join(CH, "00_front_matter.md"), encoding="utf-8").read()
    fm = fm.replace("# Front Matter\n", "")
    for b in parse_blocks(fm):
        if b[0] == "h2":
            para(doc, b[1], align="center", size=14, bold=True, space_after=8)
        elif b[0] == "para":
            para(doc, b[1], align="justify", space_after=6)
        elif b[0] == "table":
            make_data_table(doc, b[1])

    # TOC
    para(doc, "Table of Contents", align="center", size=14, bold=True, space_after=10)
    toc_p = doc.add_paragraph()
    add_field(toc_p, 'TOC \\o "1-3" \\h \\z \\u',
              "Right-click the table and choose Update Field to populate the table of contents.")

    # ----- section 3: main content
    sec3 = doc.add_section(WD_SECTION.NEW_PAGE)
    sec3.page_width, sec3.page_height = Cm(21.0), Cm(29.7)
    sec3.top_margin = sec3.bottom_margin = Cm(2.5)
    sec3.left_margin = sec3.right_margin = Cm(2.5)
    sec3.header.is_linked_to_previous = False
    sec3.footer.is_linked_to_previous = False
    hp = sec3.header.paragraphs[0]; hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hr = hp.add_run("CAPSM PhD Thesis"); hr.font.size = Pt(10); hr.font.name = "Times New Roman"
    fp = sec3.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = fp.add_run("Page "); fr.font.size = Pt(10); fr.font.name = "Times New Roman"
    add_field(fp, "PAGE")
    for r in fp.runs: r.font.size = Pt(10)

    chapter_files = ["ch01.md", "ch02.md", "ch03.md", "ch04.md", "ch05.md", "ch06.md",
                     "ch07.md", "ch08.md", "ch09.md", "ch10.md", "ch11.md", "ch12.md",
                     "ch13.md", "91_appendix_a.md", "92_appendix_b.md"]
    for fname in chapter_files:
        path = os.path.join(CH, fname)
        md = open(path, encoding="utf-8").read()
        for b in parse_blocks(md):
            kind = b[0]
            if kind == "h1":
                doc.add_heading(b[1], level=1)
            elif kind == "h2":
                doc.add_heading(b[1], level=2)
            elif kind == "h3":
                doc.add_heading(b[1], level=3)
            elif kind == "h4":
                para(doc, b[1], align="left", bold=True, space_after=4)
            elif kind == "para":
                para(doc, b[1], align="justify")
            elif kind == "li":
                para(doc, b[1], align="left", style="List Bullet", space_after=3)
            elif kind == "caption":
                para(doc, b[1], align="left", size=10, space_after=3)
            elif kind == "fig":
                cap, rel = b[1], b[2]
                img = os.path.join(FIG, rel)
                if os.path.exists(img):
                    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p.paragraph_format.space_after = Pt(2); p.paragraph_format.space_before = Pt(6)
                    with Image.open(img) as im:
                        w, h = im.size
                    disp_w = min(Inches(6.1), Emu(int(w / 300 * 914400)))
                    p.add_run().add_picture(img, width=disp_w)
                    para(doc, cap, align="center", size=10, space_after=10)
                else:
                    para(doc, f"[missing figure: {rel}]", align="center", size=10)
            elif kind == "eq":
                num, latex = b[1], b[2]
                tag = num.replace(".", "_")
                png = os.path.join(EQ, f"eq_{tag}.png")
                t = doc.add_table(rows=1, cols=2)
                t.alignment = WD_TABLE_ALIGNMENT.CENTER
                t.autofit = False
                t.columns[0].width = Cm(14.6); t.columns[1].width = Cm(1.9)
                c0, c1 = t.cell(0, 0), t.cell(0, 1)
                c0.width = Cm(14.6); c1.width = Cm(1.9)
                p0 = c0.paragraphs[0]; p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p0.paragraph_format.space_after = Pt(6); p0.paragraph_format.line_spacing = 1.0
                if os.path.exists(png):
                    with Image.open(png) as im:
                        w, h = im.size
                    disp_w = min(Inches(5.6), Emu(int(w / 400 * 914400)))
                    p0.add_run().add_picture(png, width=disp_w)
                else:
                    add_rich_runs(p0, latex, size=11, italic=True)
                p1 = c1.paragraphs[0]; p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                r1 = p1.add_run(f"({num})"); r1.font.size = Pt(11); r1.font.name = "Times New Roman"
            elif kind == "table":
                make_data_table(doc, b[1])
            elif kind == "code":
                t = doc.add_table(rows=1, cols=1)
                t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
                cell = t.cell(0, 0); set_cell_bg(cell, "F5F5F5")
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(2); p.paragraph_format.line_spacing = 1.0
                for k, line in enumerate(b[1].split("\n")):
                    if k > 0:
                        p = cell.add_paragraph()
                        p.paragraph_format.space_after = Pt(2); p.paragraph_format.line_spacing = 1.0
                    r = p.add_run(line if line else " ")
                    r.font.name = "Consolas"; r.font.size = Pt(8.5)
                    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
                para(doc, "", space_after=4, align="left")

    # ----- references
    doc.add_heading("References", level=1)
    para(doc, "All in-text citations refer to the entries below. Web sources were accessed in "
              "September 2026 during thesis preparation; dataset versions are fixed by the "
              "provenance manifest in Appendix A.", align="justify", space_after=8)
    facts = open(FACTS, encoding="utf-8").read()
    m = re.search(r"## 8\. CANONICAL REFERENCE LIST.*?\n(.*?)\n## 9\.", facts, re.S)
    refs_block = m.group(1)
    for line in refs_block.split("\n"):
        line = line.strip()
        if not re.match(r"^\[\d+\]", line): continue
        p = para(doc, "", align="left", size=10, space_after=3, hanging=1.0)
        mm = re.match(r"^(\[\d+\])\s*(.*)$", line)
        r = p.add_run(mm.group(1) + " "); r.font.size = Pt(10); r.font.name = "Times New Roman"
        rest = mm.group(2)
        url_m = re.search(r"(https?://\S+?)(?:\s|$)", rest)
        if url_m:
            before = rest[:url_m.start()]; url = url_m.group(1).rstrip(".")
            after = rest[url_m.end():]
            r2 = p.add_run(before + " "); r2.font.size = Pt(10); r2.font.name = "Times New Roman"
            add_hyperlink(p, url, url, font_size=10)
            if after.strip():
                r3 = p.add_run(" " + after.strip()); r3.font.size = Pt(10); r3.font.name = "Times New Roman"
        else:
            r2 = p.add_run(rest); r2.font.size = Pt(10); r2.font.name = "Times New Roman"

    doc.save(OUT)
    print("saved", OUT)

if __name__ == "__main__":
    build()
